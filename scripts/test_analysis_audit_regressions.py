"""Offline regressions: quantities, evidence identity, grouping and export totals."""
import re
import asyncio
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from app.schemas.ai import ExtractedDeviceSchema
from app.services.ai.analysis_pipeline_service import AnalysisPipelineService as Pipeline
from app.services.export.dynamic_grouping import DynamicGroupingEngine as Grouping
from app.services.export.quotation_exporter import QuotationExporterService as Exporter


def device(**updates):
    values = dict(category="MCB", name="MCB", spec="16A", tag="QF1", quantity=1)
    values.update(updates)
    return ExtractedDeviceSchema(**values)


class AnalysisAuditTests(unittest.TestCase):
    def test_generation_normalization_keeps_evidence_and_quantity_basis(self):
        original = device(
            box_2d=[100, 100, 200, 200], source_filename="source.jpg",
            source_type="image", source_page=1, drawing_quantity=1,
            procurement_quantity=3, quantity_basis="3XCT", quantity_confidence=0.9,
            panel_evidence_image="preview", evidence_region={"coordinate_space": "normalized_1000"},
        ).model_dump()
        normalized = []
        class StopBeforeCatalog(Exception):
            pass
        async def progress(event):
            if event.get("stage") == "catalog":
                raise StopBeforeCatalog()
        class CapturingSchema(ExtractedDeviceSchema):
            def __init__(self, **values):
                super().__init__(**values)
                normalized.append(self)
        with patch("app.services.ai.analysis_pipeline_service.ExtractedDeviceSchema", CapturingSchema):
            with self.assertRaises(StopBeforeCatalog):
                asyncio.run(Pipeline.generate_cad_and_quotation(
                    project=None, db=None, devices=[original], progress_callback=progress,
                ))
        for key in ("box_2d", "source_filename", "source_type", "source_page", "drawing_quantity",
                    "procurement_quantity", "quantity_basis", "quantity_confidence",
                    "panel_evidence_image", "evidence_region"):
            self.assertEqual(getattr(normalized[0], key), original[key], key)

    def test_explicit_fuse_quantity_wins_over_three_phase_heuristic(self):
        dev = device(category="FUSE", name="Cầu chì x6", spec="3 pha")
        Pipeline._infer_practical_quantities([dev])
        Pipeline._infer_practical_quantities([dev])
        self.assertEqual(dev.quantity, 6)
        self.assertEqual(dev.procurement_quantity, 6)

    def test_three_ct_quantity_is_idempotent(self):
        dev = device(category="CT", name="3XCT")
        for _ in range(2):
            Pipeline._infer_practical_quantities([dev])
        self.assertEqual((dev.drawing_quantity, dev.quantity), (1, 3))

    def test_identical_tags_in_different_panels_keep_distinct_boxes(self):
        devices = [device(panel_code="DB1"), device(panel_code="DB2")]
        boxes = [[100, 100, 200, 200], [600, 600, 700, 700]]
        Pipeline._apply_verified_boxes(devices, {"boxes": [
            {"device_id": str(i), "tag": "QF1", "verified": True, "box_2d": box}
            for i, box in enumerate(boxes)
        ]})
        self.assertEqual([d.box_2d for d in devices], boxes)

    def test_ambiguous_or_invalid_verifier_output_does_not_invent_box(self):
        for candidates in [
            [{"tag": "QF1", "verified": True, "box_2d": [1, 1, 2, 2]}],
            [{"device_id": "0", "verified": True, "box_2d": [9, 9, 1, 1]}],
            [{"device_id": "0", "verified": True, "box_2d": [1, 1, 2, 2]}] * 2,
        ]:
            dev = device(box_2d=[100, 100, 200, 200])
            Pipeline._apply_verified_boxes([dev], {"boxes": candidates})
            self.assertIsNone(dev.box_2d)

    def test_crop_stays_in_image_and_invalid_box_is_rejected(self):
        crop = Pipeline._evidence_crop_pixels([0, 0, 100, 100], 1000, 500)
        self.assertEqual(crop[:2], (0, 0))
        self.assertLessEqual(crop[2], 1000)
        self.assertLessEqual(crop[3], 500)
        self.assertIsNone(Pipeline._evidence_crop_pixels([200, 200, 100, 100], 1000, 500))

    def test_grouping_preserves_panel_and_breaking_capacity(self):
        for changes in [dict(panel_code="DB2"), dict(icu_ka=10)]:
            first = device(panel_code="DB1", icu_ka=6).model_dump()
            second = dict(first, **changes)
            self.assertEqual(len(Grouping.merge_duplicates([first, second])), 2)

    def test_grouping_missing_quantity_defaults_and_sums(self):
        raw = {"category": "MCB", "name": "MCB", "spec": "16A"}
        self.assertEqual(Grouping.merge_duplicates([raw, raw])[0]["quantity"], 2)


class QuotationFormulaTests(unittest.TestCase):
    def test_multi_panel_formulas_have_no_cycle_and_correct_totals(self):
        rows = [
            {"row_type": "panel_header", "name": "DB1", "quantity": 2},
            {"row_type": "item", "name": "MCB", "quantity": 3, "unit_price": 100},
            {"row_type": "accessory", "name": "AUX", "quantity": 1, "unit_price": 50},
            {"row_type": "panel_header", "name": "DB2", "quantity": 1},
            {"row_type": "item", "name": "MCCB", "quantity": 2, "unit_price": 200},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            with patch("app.services.export.quotation_exporter.settings.EXPORT_DIR", tmp):
                path = Exporter.export(rows, project_name="Audit", vat_percent=10)
            wb = load_workbook(path)
            ws = wb.active

            # Evaluate this exporter's scalar arithmetic and SUM formulas;
            # fail explicitly on a cycle instead of relying on Excel caches.
            def evaluate(address, stack=()):
                self.assertNotIn(address, stack, "Circular formula reference")
                value = ws[address].value
                if not isinstance(value, str) or not value.startswith("="):
                    return float(value or 0)
                expression = value[1:]
                def sum_args(match):
                    total = 0
                    for arg in match.group(1).split(","):
                        if ":" in arg:
                            start, end = arg.split(":")
                            for row in ws[f"{start}:{end}"]:
                                for cell in row:
                                    total += evaluate(cell.coordinate, stack + (address,))
                        else:
                            total += evaluate(arg, stack + (address,))
                    return str(total)
                expression = re.sub(r"SUM\(([^)]+)\)", sum_args, expression)
                expression = re.sub(r"\b[A-Z]+\d+\b", lambda m: str(evaluate(m[0], stack + (address,))), expression)
                self.assertRegex(expression, r"^[0-9. +*/()-]+$")
                return eval(expression, {"__builtins__": {}}, {})

            self.assertEqual(evaluate("H2"), 700)
            self.assertEqual(evaluate("H5"), 400)
            self.assertEqual(evaluate("H7"), 1100)
            self.assertEqual(evaluate("H8"), 110)
            self.assertEqual(evaluate("H9"), 1210)
            wb.close()

    def test_empty_quote_does_not_reference_itself(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch("app.services.export.quotation_exporter.settings.EXPORT_DIR", tmp):
                path = Exporter.export([{"row_type": "panel_header", "name": "Empty"}])
            wb = load_workbook(path)
            self.assertEqual(wb.active["H3"].value, "=0")
            wb.close()


if __name__ == "__main__":
    unittest.main()
