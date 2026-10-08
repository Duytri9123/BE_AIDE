"""Run the real CAD/quotation service with isolated SQL state and supplied devices."""
import asyncio,json,sys
from pathlib import Path
from sqlalchemy.ext.asyncio import create_async_engine,async_sessionmaker
from sqlalchemy import select
from app.models.base import Base
from app.models.project import Project
from app.models.project_file import ProjectFile
from app.services.ai.analysis_pipeline_service import AnalysisPipelineService
import app.models
async def main():
 devices=json.loads(Path(sys.argv[1]).read_text(encoding='utf-8'))
 engine=create_async_engine('sqlite+aiosqlite:///:memory:')
 async with engine.begin() as conn:await conn.run_sync(Base.metadata.create_all)
 async with async_sessionmaker(engine,expire_on_commit=False)() as db:
  project=Project(id=990008,user_id=1,name='Integration CAD TDT');db.add(project);await db.commit()
  events=[]
  result=await AnalysisPipelineService.generate_cad_and_quotation(project,db,devices,enclosure_dimensions='1000x600x300',panel_code='TĐT',per_panel=True,progress_callback=events.append)
  assert result['cad_file'],events
  assert result['quotation_file'],events
  assert result['cad_layout']['branch_arrangement']=='two_vertical_banks'
  rows=[r for r in result['quotation_rows'] if r.get('id')=='layout-neutral']
  assert len(rows)==1 and rows[0]['unit_price'] is None,rows
  from openpyxl import load_workbook
  workbook=load_workbook(result['quotation_file']['file_path'])
  descriptions=[str(row[1].value or '') for row in workbook.active.iter_rows() if len(row)>1]
  neutral=next(p for p in result['cad_layout']['placements'] if p.get('status')=='custom_fabricated_review')
  expected_length=f"{neutral['length_mm']:g} mm"
  assert any('Thanh đồng N gia công' in text and expected_length in text for text in descriptions), 'Neutral layout length missing from exported file'
  workbook.close()
  files=(await db.execute(select(ProjectFile))).scalars().all()
  assert len(files)==2
  for file in files:assert Path(file.file_path).is_file()
  out=Path('tmp/system_cad_integration.json');out.parent.mkdir(exist_ok=True);out.write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
  print(json.dumps({'status':'passed','files':[f.filename for f in files],'neutral':rows,'report':str(out)},ensure_ascii=False))
 await engine.dispose()
if __name__=='__main__':asyncio.run(main())
