import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
import ezdxf
from openpyxl import load_workbook
from pydantic import ValidationError
from app.services.ai.system_completeness import review_system, technical_audit
from app.services.export.quotation_exporter import QuotationExporterService
from app.api.v1.endpoints.export import ExcelExportPayload
from app.api.v1.endpoints.cad_library import LIBRARY, list_layouts, download_layout


class ClusterTests(unittest.TestCase):
    def test_untagged_ct_keeps_explicit_quantity_and_parent(self):
        from app.schemas.ai import ExtractedDeviceSchema
        from app.services.ai.analysis_pipeline_service import AnalysisPipelineService as P
        parent = ExtractedDeviceSchema(category="MCCB", name="CB", spec="63A", tag="QF1",
            accompanying_accessories=[dict(name="CT", quantity=1, evidence="Ký hiệu CT cạnh QF1")])
        result = P._promote_embedded_accessories_to_devices([parent])
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0].accompanying_accessories[0]["quantity"], 1)

    def test_tagged_components_in_distinct_clusters_are_not_dropped(self):
        from app.schemas.ai import ExtractedDeviceSchema
        from app.services.ai.analysis_pipeline_service import AnalysisPipelineService as P
        parents = [ExtractedDeviceSchema(category="MCCB", name="CB", spec="63A", tag=f"QF{i}",
            accompanying_accessories=[dict(tag=f"CT{i}", category="CT", name="CT", quantity=1, evidence=f"CT{i}")]) for i in (1, 2)]
        result = P._promote_embedded_accessories_to_devices(parents)
        self.assertEqual(len(result), 4)
        self.assertEqual([d.tag for d in result[2:]], ["CT1", "CT2"])
        self.assertTrue(all(d.quantity == 1 for d in result[2:]))

    def test_unverified_components_are_not_added_to_bom(self):
        devices = [dict(name="Đồng hồ và chuyển mạch AS", category="METER", panel_code="DB1", tag="PA",
                        accompanying_accessories=[dict(name="AS", evidence="Nhãn AS cạnh đồng hồ")],
                        inferred_components=[dict(name="Cầu chì", quantity=3)])]
        before = json.dumps(devices)
        result = review_system(devices)
        self.assertEqual(json.dumps(devices), before)
        self.assertEqual(result["automatically_added_devices"], 0)
        self.assertFalse(result["release_ready"])
        self.assertEqual(len(result["clusters"][0]["observed_components"]), 1)
        self.assertEqual(len(result["clusters"][0]["review_components"]), 1)

    def test_upstream_reference_does_not_match_other_panel(self):
        result = review_system([
            dict(tag="QF1", panel_code="DB1", name="CB"),
            dict(tag="KM1", panel_code="DB2", name="Contactor", upstream_device="QF1"),
        ])
        self.assertTrue(any(i["title"] == "Chưa tìm thấy thiết bị cấp nguồn" for i in result["issues"]))

    def test_no_fake_safety_score(self):
        audit = technical_audit([dict(name="MCB", category="MCB", in_a=16, poles=1, icu_ka=6)])
        self.assertIsNone(audit["overall_score"])
        self.assertFalse(audit["release_ready"])
        self.assertTrue(all(i["status"] != "PASS" for i in audit["protection_coordination"]))


class DiscountTests(unittest.TestCase):
    def test_validate_rate(self):
        for value in (-1, 101, float("nan"), float("inf")):
            with self.assertRaises(ValidationError):
                ExcelExportPayload(project_id=1, manufacturer_discounts={"LS": value})

    def test_export_keeps_list_price_and_live_net_formula(self):
        rows = [dict(row_type="panel_header", name="DB"),
                dict(row_type="item", name="MCB", origin="LS", quantity=2, unit_price=100),
                dict(row_type="item", name="Vỏ", origin="VN", quantity=1, unit_price=50)]
        with tempfile.TemporaryDirectory() as tmp:
            with patch("app.services.export.quotation_exporter.settings.EXPORT_DIR", tmp):
                path = QuotationExporterService.export(rows, manufacturer_discounts={"ls": 25})
            wb = load_workbook(path)
            ws = wb.active
            self.assertEqual(ws["J3"].value, 100)
            self.assertEqual(ws["K3"].value, 25)
            self.assertEqual(ws["G3"].value, "=J3*(1-K3/100)")
            self.assertEqual(ws["H3"].value, "=F3*G3")
            self.assertEqual(ws["G4"].value, "=J4*(1-K4/100)")
            self.assertEqual(ws["J4"].value, 50)
            self.assertEqual(ws["K4"].value, 0)
            wb.close()


class LibraryTests(unittest.TestCase):
    def test_svg_preview_is_valid_and_unknown_asset_is_rejected(self):
        import xml.etree.ElementTree as ET
        from fastapi import HTTPException
        from app.api.v1.endpoints.cad_library import render_layout_svg
        root = ET.fromstring(render_layout_svg("7a0ea9fb9719bdce4a3b"))
        self.assertTrue(root.tag.endswith("svg"))
        with self.assertRaises(HTTPException) as error:
            render_layout_svg("unknown-asset")
        self.assertEqual(error.exception.status_code, 404)

    def test_library_endpoint_matches_manifest_and_filters(self):
        items = list_layouts("800AF 3P")["items"]
        self.assertTrue(items)
        for item in items:
            response = download_layout(item["id"])
            self.assertTrue(Path(response.path).is_file())
            self.assertFalse(item["model_verified"])

    def test_representative_layout_is_mm_and_at_origin(self):
        from ezdxf import bbox
        item = next(i for i in list_layouts()["items"] if i["name"] == "LS 800AF 3P")
        doc = ezdxf.readfile(LIBRARY / item["filename"])
        bounds = bbox.extents(doc.modelspace())
        self.assertEqual(doc.units, 4)
        self.assertAlmostEqual(bounds.extmin.x, 0, places=2)
        self.assertAlmostEqual(bounds.extmin.y, 0, places=2)
        self.assertFalse(doc.audit().errors)


if __name__ == "__main__":
    unittest.main()
