import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from app.services.ai.procurement_review import build_quotation_rows, quotation_section
from app.services.export.quotation_exporter import QuotationExporterService


class ReferenceLayoutTests(unittest.TestCase):
    def test_saved_commercial_draft_has_three_sections_and_no_design_cable_line(self):
        from app.services.export.quotation_names import quotation_display_rows
        rows=[dict(id='P',row_type='panel_header',name='P'),
            dict(row_type='section_header',name='Main'),
            dict(id='main',row_type='item',commercial_display=True,category='MCCB',name='MCCB',quantity=1),
            dict(row_type='section_header',name='Metering'),
            dict(id='ct',row_type='item',commercial_display=True,category='CT',name='CT',quantity=3),
            dict(id='cable',row_type='item',commercial_display=True,category='CABLE',name='Cáp nguồn',quantity=1),
            dict(row_type='section_header',name='Outgoing'),
            dict(id='branch',row_type='item',commercial_display=True,category='MCB',name='MCB',quantity=4)]
        displayed=quotation_display_rows(rows)
        self.assertEqual([r['name'] for r in displayed if r['row_type']=='section_header'],['Đầu vào','Đầu ra'])
        self.assertEqual([r['id'] for r in displayed if r['row_type']=='item'],['main','ct','branch'])
        self.assertEqual(rows[1]['name'],'Main')

    def test_explicit_branch_names_group_as_outgoing_without_guessing_current(self):
        self.assertEqual(quotation_section({'name':'MCCB nhánh WP1 đến WP4','in_a':630}), 'Đầu ra')
        self.assertEqual(quotation_section({'name':'MCCB dự phòng đáy'}), 'Đầu ra')
        self.assertEqual(quotation_section({'name':'MCCB nhánh','section':'Đầu vào'}), 'Đầu vào')
        self.assertEqual(quotation_section({'name':'Đồng hồ','category':'METER'}), 'Đầu vào')
    def test_source_dimensions_and_unselected_price_options_are_visible_without_firm_price(self):
        device=dict(category='MCCB',name='Main',spec='3P 630A',poles=3,in_a=630,
                    brand='CHINT',panel_code='MDB101',quantity=1,
                    pricebook_match={'status':'needs_selection','candidates':[
                        {'sku':'NXM-630S','brand':'CHINT','price':7526000}]})
        rows=build_quotation_rows([device],{},panels=[{'panel_code':'MDB101','dimension':'1800x800x450 x1.5mm'}],refresh_prices=False)
        enclosure=next(r for r in rows if r['id']=='MDB101-assembly-ENCLOSURE')
        self.assertIn('1800x800x450',enclosure['spec'])
        item=next(r for r in rows if r['row_type']=='item' and not r.get('commercial_package'))
        self.assertEqual(item['notes'],'')
        self.assertEqual(item['pricebook_match']['candidates'][0]['sku'],'NXM-630S')
        self.assertIsNone(item['unit_price'])
        self.assertEqual(item['sku'],'')
    def test_panels_and_sections_keep_every_device_and_accessory(self):
        devices = [dict(category='MCCB', name='Feeder', spec='3P160A', section='Đầu ra',
                        panel_code='MDB101', quantity=4),
                   dict(category='MCCB', name='Main', spec='3P630A', section='Đầu vào',
                        panel_code='MDB101', quantity=1,
                        accompanying_accessories=[dict(name='Cuộn cắt', quantity=1, evidence='Nguồn ghi')]),
                   dict(category='CT', name='CT', spec='600/5A', section='Đầu vào',
                        panel_code='MDB102', quantity=3)]
        with patch('app.services.ai.procurement_review.assembly_requirements', return_value=[]):
            rows = build_quotation_rows(devices, {}, refresh_prices=False)
        self.assertEqual([r['name'] for r in rows if r['row_type'] == 'section_header'],
                         ['Vỏ tủ + phụ kiện', 'Đầu vào', 'Đầu ra', 'Vỏ tủ + phụ kiện', 'Đầu vào'])
        self.assertEqual([r['quantity'] for r in rows if r['row_type'] == 'item' and not r.get('commercial_package')], [1, 4, 3])
        self.assertEqual(len([r for r in rows if r.get('commercial_package')]),8)
        accessory = next(r for r in rows if r['row_type'] == 'accessory')
        parent = rows[rows.index(accessory) - 1]
        self.assertEqual(accessory['parent_id'], parent['id'])
        self.assertEqual([r['tt'] for r in rows if r['row_type'] == 'panel_header'], ['1', '2'])
        self.assertTrue(all(r['unit_price'] is None for r in rows if r['row_type'] == 'item'))

    def test_export_nine_visible_columns_and_formula_inputs_preserved(self):
        rows = [dict(row_type='panel_header', name='MDB101', quantity=1),
                dict(row_type='section_header', name='Đầu vào'),
                dict(row_type='item', name='MCCB', quantity=1, unit_price=100, origin='CHINT',
                     notes='Thông tin xác minh mã hàng và cấu hình cần được hiển thị đầy đủ trong bảng báo giá. '*3)]
        with tempfile.TemporaryDirectory() as folder, patch('app.services.export.quotation_exporter.settings.EXPORT_DIR', folder):
            book = load_workbook(QuotationExporterService.export(rows, proposals=[]))
            sheet = book.active
            self.assertEqual(sheet['B2'].value, 'TỦ ĐIỆN MDB101')
            self.assertEqual(sheet['B3'].value, 'Đầu vào')
            self.assertEqual(sheet['G4'].value, '=J4*(1-K4/100)')
            self.assertEqual(sheet['J4'].value, 100)
            self.assertGreater(sheet.row_dimensions[4].height, 24)
            self.assertTrue(sheet.column_dimensions['J'].hidden)
            self.assertTrue(sheet.column_dimensions['K'].hidden)
            self.assertIn('$A$1:$I$', str(sheet.print_area))
            book.close()

    def test_compact_export_hides_notes_and_leaves_discount_to_customer(self):
        rows=[dict(row_type='panel_header',name='MDB',quantity=1,notes='Tên thiết kế nội bộ'),
              dict(row_type='item',commercial_display=True,name='MCCB',spec='3P 630A',origin='CHINT',quantity=2,unit_price=100,notes='')]
        with tempfile.TemporaryDirectory() as folder, patch('app.services.export.quotation_exporter.settings.EXPORT_DIR',folder):
            for discounts,expected in (({},0),({'CHINT':51},51)):
                book=load_workbook(QuotationExporterService.export(rows,manufacturer_discounts=discounts,
                    proposals=[{'recommended_model':'Không đưa phân tích thiết kế vào báo giá'}]))
                sheet=book.active
                self.assertEqual(sheet['B3'].value,'MCCB 3P 630A')
                self.assertTrue(sheet.column_dimensions['I'].hidden)
                self.assertIn('$A$1:$H$',str(sheet.print_area))
                self.assertEqual(sheet['K3'].value,expected)
                self.assertFalse(any('Không đưa phân tích' in str(cell.value) for row in sheet for cell in row))
                book.close()


if __name__ == '__main__':
    unittest.main()
