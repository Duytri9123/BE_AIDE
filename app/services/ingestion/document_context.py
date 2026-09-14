"""Extract bounded, source-attributed context from heterogeneous project files.

This module deliberately does not decide whether a file is a drawing.  It makes
all readable attachments available to the orchestration/AI layer; the role of a
file is decided against the current user request.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, Iterable, List


class DocumentContextService:
    MAX_FILE_CHARS = 12_000
    MAX_CELL_ROWS = 250

    @staticmethod
    def _bounded(value: Any, limit: int | None = None) -> str:
        text = str(value or "").replace("\x00", " ").strip()
        return text[: limit or DocumentContextService.MAX_FILE_CHARS]

    @staticmethod
    def _read_text(path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeError:
                continue
        return ""

    @staticmethod
    def _read_pdf(path: Path) -> str:
        import pdfplumber

        parts: List[str] = []
        with pdfplumber.open(str(path)) as document:
            for page_no, page in enumerate(document.pages, 1):
                text = (page.extract_text() or "").strip()
                if text:
                    parts.append(f"[Trang {page_no}]\n{text}")
                if sum(len(part) for part in parts) >= DocumentContextService.MAX_FILE_CHARS:
                    break
        return "\n\n".join(parts)

    @staticmethod
    def _read_docx(path: Path) -> str:
        from docx import Document

        doc = Document(str(path))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                if any(values):
                    parts.append(" | ".join(values))
        return "\n".join(parts)

    @staticmethod
    def _read_xlsx(path: Path) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=False)
        parts: List[str] = []
        try:
            for sheet in workbook.worksheets:
                parts.append(f"[Sheet: {sheet.title}]")
                for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    if row_no > DocumentContextService.MAX_CELL_ROWS:
                        break
                    values = [str(value).strip() for value in row if value not in (None, "")]
                    if values:
                        parts.append(" | ".join(values))
                    if sum(len(part) for part in parts) >= DocumentContextService.MAX_FILE_CHARS:
                        break
        finally:
            workbook.close()
        return "\n".join(parts)

    @staticmethod
    def _read_delimited(path: Path) -> str:
        text = DocumentContextService._read_text(path)
        rows: List[str] = []
        try:
            dialect = csv.Sniffer().sniff(text[:4096])
            reader: Iterable[List[str]] = csv.reader(text.splitlines(), dialect)
        except csv.Error:
            reader = csv.reader(text.splitlines())
        for row_no, row in enumerate(reader, 1):
            if row_no > DocumentContextService.MAX_CELL_ROWS:
                break
            rows.append(" | ".join(str(value).strip() for value in row))
        return "\n".join(rows)

    @staticmethod
    def extract(file_obj: Any) -> Dict[str, Any]:
        filename = str(getattr(file_obj, "filename", "") or "")
        path = Path(str(getattr(file_obj, "file_path", "") or ""))
        ext = path.suffix.lower().lstrip(".") or filename.rsplit(".", 1)[-1].lower()
        result: Dict[str, Any] = {
            "file_id": getattr(file_obj, "id", None),
            "filename": filename,
            "extension": ext,
            "mime_type": str(getattr(file_obj, "file_type", "") or ""),
            "exists": path.is_file(),
            "text": "",
            "read_error": "",
        }
        if not path.is_file():
            result["read_error"] = "file_not_found"
            return result

        try:
            if ext == "pdf":
                text = DocumentContextService._read_pdf(path)
            elif ext == "docx":
                text = DocumentContextService._read_docx(path)
            elif ext in {"xlsx", "xlsm"}:
                text = DocumentContextService._read_xlsx(path)
            elif ext in {"csv", "tsv"}:
                text = DocumentContextService._read_delimited(path)
            elif ext == "json":
                parsed = json.loads(DocumentContextService._read_text(path))
                text = json.dumps(parsed, ensure_ascii=False, indent=2)
            elif ext in {"txt", "md", "log", "xml", "html", "yaml", "yml"}:
                text = DocumentContextService._read_text(path)
            else:
                text = ""
            result["text"] = DocumentContextService._bounded(text)
        except Exception as exc:  # optional readers must not discard the attachment
            result["read_error"] = f"{type(exc).__name__}: {exc}"[:500]
        return result

    @staticmethod
    def build_prompt_context(contexts: List[Dict[str, Any]], max_total_chars: int = 30_000) -> str:
        blocks: List[str] = []
        used = 0
        for context in contexts:
            text = str(context.get("text") or "").strip()
            if not text:
                continue
            header = f"--- TỆP THAM CHIẾU: {context.get('filename') or 'không tên'} ---\n"
            remaining = max_total_chars - used - len(header)
            if remaining <= 0:
                break
            block = header + text[:remaining]
            blocks.append(block)
            used += len(block)
        return "\n\n".join(blocks)
