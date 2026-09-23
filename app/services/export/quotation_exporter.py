import os
import uuid
from datetime import datetime
from pathlib import Path
from typing import Optional, List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from app.services.export.dynamic_grouping import DynamicGroupingEngine
from app.core.config import settings
from app.core.constants import BRAND_CATALOG_MAPPING

class QuotationExporterService:
    @staticmethod
    def export(
        devices: list,
        project_name: str = "Tủ Điện",
        brand_preference: str = "Theo thiết kế",
        proposals: Optional[list] = None,
        vat_percent: Optional[float] = None,
        filename_suffix: Optional[str] = None,
        manufacturer_discounts: Optional[dict] = None,
    ) -> str:
        """Xuất bảng báo giá thiết bị tủ điện chuẩn kỹ thuật ra file Excel (.xlsx)."""
        wb = Workbook()
        discounts = {}
        import math
        for brand, value in (manufacturer_discounts or {}).items():
            rate = float(value)
            if not math.isfinite(rate) or not 0 <= rate <= 100:
                raise ValueError("Chiết khấu hãng phải từ 0 đến 100%")
            discounts[str(brand).strip().casefold()] = rate
        ws = wb.active
        ws.title = "Bảng Báo Giá"
        ws.views.sheetView[0].showGridLines = True

        # Styles definition (Times New Roman)
        font_header = Font(name="Times New Roman", size=10, bold=True, color="000000")
        font_data = Font(name="Times New Roman", size=10, color="000000")
        font_data_bold = Font(name="Times New Roman", size=10, bold=True, color="000000")
        font_panel = Font(name="Times New Roman", size=10, bold=True, color="000000")
        font_section = Font(name="Times New Roman", size=10, bold=True, color="000000")
        font_grand_total = Font(name="Times New Roman", size=11, bold=True, color="000000")

        fill_header = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid") # Gold Yellow
        fill_panel = PatternFill(start_color="A9D08E", end_color="A9D08E", fill_type="solid")   # Light Green
        fill_section = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid") # Light Gray
        fill_total = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")   # Gold Yellow

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
        align_right = Alignment(horizontal="right", vertical="center")

        thin_side = Side(style="thin", color="000000")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        # 1. Table Headers (Gold Yellow #FFC000)
        headers = [
            "TT",
            "MÔ TẢ CHI TIẾT",
            "MÃ HÀNG",
            "XUẤT XỨ",
            "ĐƠN VỊ",
            "SỐ LƯỢNG",
            "ĐƠN GIÁ",
            "THÀNH TIỀN",
            "GHI CHÚ",
            "GIÁ NIÊM YẾT",
            "CHIẾT KHẤU HÃNG (%)"
        ]

        header_row = 1
        for col_idx, h_text in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = border_all
            ws.row_dimensions[header_row].height = 28

        current_row = 2
        item_rows_indices = []
        panel_item_rows = {}

        def write_section(title: str):
            nonlocal current_row
            sec_cell_tt = ws.cell(row=current_row, column=1, value="*")
            sec_cell_tt.font = font_section
            sec_cell_tt.fill = fill_section
            sec_cell_tt.alignment = align_center
            sec_cell_tt.border = border_all

            sec_cell_title = ws.cell(row=current_row, column=2, value=title)
            sec_cell_title.font = font_section
            sec_cell_title.fill = fill_section
            sec_cell_title.alignment = align_left
            sec_cell_title.border = border_all

            for c in range(3, 10):
                empty_c = ws.cell(row=current_row, column=c, value="")
                empty_c.fill = fill_section
                empty_c.border = border_all
            ws.row_dimensions[current_row].height = 22
            current_row += 1

        def write_item(name: str, sku: str, origin: str, unit: str, qty: float, price: int, notes: str = "", tt: str = "+"):
            nonlocal current_row
            discount = discounts.get(origin.strip().casefold(), 0)
            net_price = f"=J{current_row}*(1-K{current_row}/100)"
            item_cells = [
                (tt, align_center, False),
                (name, align_left, False),
                (sku, align_center, False),
                (origin, align_center, False),
                (unit, align_center, False),
                (qty, align_center, False),
                (net_price, align_right, False),
                (f"=F{current_row}*G{current_row}", align_right, True),
                (notes, align_left, False),
                (price, align_right, False),
                (discount, align_right, False),
            ]
            for col_idx, (val, al, is_b) in enumerate(item_cells, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = font_data_bold if is_b else font_data
                cell.alignment = al
                cell.border = border_all
                if col_idx in (7, 8):
                    cell.number_format = '#,##0'
                if col_idx == 6:
                    cell.number_format = '#,##0.00'
            ws.row_dimensions[current_row].height = 24
            item_rows_indices.append(current_row)
            panel_item_rows.setdefault(panel_row_idx, []).append(current_row)
            current_row += 1

        # Check if devices is already a full structured spreadsheet rows list
        has_row_types = any(isinstance(d, dict) and d.get("row_type") for d in devices)

        panel_row_idx = 2
        clean_p_name = project_name.upper().replace("DỰ ÁN", "").strip() or "TSXĐ-630A"

        if has_row_types:
            for r in devices:
                r_type = r.get("row_type")
                if r_type == "panel_header":
                    p_name = str(r.get("name") or f"TỦ ĐIỆN {clean_p_name}")
                    p_origin = str(r.get("origin") or "VN")
                    p_unit = str(r.get("unit") or "Tủ")
                    p_qty = float(r.get("quantity") or 1.0)
                    panel_cells = [
                        (str(r.get("tt") or "1"), align_center, font_panel),
                        (p_name, align_left, font_panel),
                        (str(r.get("sku") or ""), align_center, font_panel),
                        (p_origin, align_center, font_panel),
                        (p_unit, align_center, font_panel),
                        (p_qty, align_center, font_panel),
                        (0, align_right, font_panel),
                        (0, align_right, font_panel),
                        (str(r.get("notes") or ""), align_left, font_panel),
                    ]
                    panel_row_idx = current_row
                    for col_idx, (val, al, fn) in enumerate(panel_cells, 1):
                        cell = ws.cell(row=current_row, column=col_idx, value=val)
                        cell.font = fn
                        cell.fill = fill_panel
                        cell.alignment = al
                        cell.border = border_all
                        if col_idx in (7, 8):
                            cell.number_format = '#,##0'
                    ws.row_dimensions[current_row].height = 24
                    current_row += 1

                elif r_type == "section_header":
                    write_section(str(r.get("name") or "Mục phân nhóm"))

                elif r_type in ["item", "accessory"]:
                    name = str(r.get("name") or "Thiết bị")
                    sku = str(r.get("sku") or "")
                    origin = str(r.get("origin") or "VN")
                    unit = str(r.get("unit") or "Cái")
                    qty = float(r.get("quantity") or 1.0)
                    price = int(r.get("unit_price") or r.get("price") or 0)
                    notes = str(r.get("notes") or "")
                    tt = str(r.get("tt") or ("↳" if r_type == "accessory" or r.get("is_accessory") else "+"))
                    write_item(name, sku, origin, unit, qty, price, notes, tt=tt)
        else:
            # Write Main Panel Row
            panel_cells = [
                ("1", align_center, font_panel),
                (f"TỦ ĐIỆN {clean_p_name}", align_left, font_panel),
                ("", align_center, font_panel),
                ("VN", align_center, font_panel),
                ("Tủ", align_center, font_panel),
                (1.0, align_center, font_panel),
                (0, align_right, font_panel),
                (0, align_right, font_panel),
                ("", align_left, font_panel),
            ]
            panel_row_idx = current_row
            for col_idx, (val, al, fn) in enumerate(panel_cells, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = fn
                cell.fill = fill_panel
                cell.alignment = al
                cell.border = border_all
                if col_idx in (7, 8):
                    cell.number_format = '#,##0'
            ws.row_dimensions[current_row].height = 24
            current_row += 1

            # Group raw devices dynamically
            sections_map = {}
            for dev in devices:
                sec = str(dev.get("section") or "").strip()
                if not sec:
                    cat = str(dev.get("category") or "").upper()
                    d_name = str(dev.get("name") or "").lower()
                    in_a = dev.get("in_a") or 0
                    if "ACB" in cat or in_a >= 400 or "incomer" in d_name or "tổng" in d_name:
                        sec = "Đầu vào"
                    elif "QUẠT" in d_name.upper() or "NHIỆT" in d_name.upper() or "LÀM MÁT" in d_name.upper():
                        sec = "Làm mát cho ngăn tủ rack"
                    elif "ĐÈN" in d_name.upper() or "ĐỒNG HỒ" in d_name.upper() or "VOLT" in d_name.upper() or "AMPE" in d_name.upper():
                        sec = "Đo lường & Giám sát"
                    elif "TỤ BÙ" in d_name.upper() or "BÙ" in d_name.upper() or "CAPACITOR" in cat:
                        sec = "Bù công suất"
                    elif "TIMER" in cat or "CONTACTOR" in cat or "RELAY" in cat:
                        sec = "Điều khiển & Phụ trợ"
                    else:
                        sec = "Đầu ra"
                if sec not in sections_map:
                    sections_map[sec] = []
                sections_map[sec].append(dev)

            # Section 1: Vỏ tủ + phụ kiện
            from app.services.bom.bom_orchestrator import BomOrchestratorService
            from app.services.bom.busbar_calculator import BusbarCalculatorService
            from app.services.device_catalog_engine import DeviceCatalogEngine

            eng = DeviceCatalogEngine.get_instance()
            bom_calc = BomOrchestratorService.calculate_full_bom("1", devices)
            enc_calc = bom_calc.enclosure
            busbar_calc = bom_calc.busbar
            labor_calc = bom_calc.labor

            enc_desc = (
                f"Vỏ tủ điện sơn tĩnh điện tiêu chuẩn công nghiệp\n"
                f"+ KT : H{enc_calc.get('H')}xW{enc_calc.get('W')}xD{enc_calc.get('D')}xT{enc_calc.get('tole_thickness_mm')}mm\n"
                f"+ Cấp bảo vệ IP54"
            )
            write_section("Vỏ tủ + phụ kiện")
            write_item(enc_desc, "Sơn tĩnh điện", "VN", "Cái", 1.0, enc_calc.get("estimated_price", 0), "Sơn tĩnh điện RAL 7035")

            incomer_a = max([float(d.get("in_a") or 0) for d in devices] or [63.0])
            need_busbar_exp = BusbarCalculatorService.needs_busbar(incomer_a, devices)
            if need_busbar_exp:
                write_item(
                    f"Đồng thanh cái mạ thiếc bọc co nhiệt (Định mức In={int(incomer_a)}A, {busbar_calc.get('section_mm2')}mm2)",
                    "Đồng đỏ 99.9%",
                    "VN",
                    "Hệ",
                    1.0,
                    busbar_calc.get("unit_price", 0),
                    f"Gia công uốn đột CNC, bọc co nhiệt (Cu {busbar_calc.get('mass_Cu_kg')}kg)"
                )
            
            acc_total = sum(float(a.get("quantity", 1)) * float(a.get("unit_price", 0)) for a in bom_calc.accessories)
            acc_name_exp = "Vật tư phụ tủ điện (đầu cosse SC động lực, dây điều khiển Cadivi VSF, máng cáp nhựa PVC, sứ đỡ thanh cái SM)" if need_busbar_exp else "Vật tư phụ tủ điện (Cầu lược 3P/1P phân phối MCB, đầu cosse ghim, dây điều khiển Cadivi VSF, máng cáp nhựa PVC, kẹp tiếp địa)"
            write_item(acc_name_exp, "Trọn gói", "VN", "Tủ", 1.0, int(acc_total))

            hourly_rate = settings.DEFAULT_LABOR_HOURLY_RATE
            labor_cost = int(float(labor_calc.get("T_total_hours", 0)) * hourly_rate)
            write_item("Nhân công lắp ráp, đấu nối & kiểm tra xuất xưởng (QC)", "QC-LABOR", "VN", "Tủ", 1.0, labor_cost)

            for sec_name, dev_list in sections_map.items():
                write_section(sec_name)
                merged_list = DynamicGroupingEngine.merge_duplicates(dev_list)
                for d in merged_list:
                    d_cat = str(d.get("category") or "")
                    d_in = d.get("in_a")
                    d_poles = d.get("poles")
                    d_price = int(d.get("unit_price") or d.get("price") or 0)
                    write_item(
                        str(d.get("name") or "Thiết bị đóng cắt"),
                        str(d.get("sku") or d.get("part_number") or d.get("spec") or ""),
                        str(d.get("brand") or brand_preference).replace(" Electric", ""),
                        str(d.get("unit") or "Cái"),
                        float(d.get("quantity") or 1.0),
                        d_price,
                        str(d.get("notes") or "")
                    )

        # 4. Summary Rows (Matching template)
        # Each panel totals only its own detail rows. Referencing the grand
        # total here creates a cycle when a later panel lies in its SUM range.
        for panel_row, detail_rows in panel_item_rows.items():
            refs = f"H{detail_rows[0]}:H{detail_rows[-1]}"
            ws.cell(row=panel_row, column=7, value=f"=SUM({refs})")
            ws.cell(row=panel_row, column=8, value=f"=F{panel_row}*G{panel_row}")

        # Row: TỔNG GIÁ TRỊ TRƯỚC THUẾ
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c_sub = ws.cell(row=current_row, column=1, value="TỔNG GIÁ TRỊ TRƯỚC THUẾ")
        c_sub.font = font_grand_total
        c_sub.alignment = align_right

        panel_refs = ",".join(f"H{row}" for row in panel_item_rows)
        tot_formula = f"=SUM({panel_refs})" if panel_refs else "=0"
        c_sub_val = ws.cell(row=current_row, column=8, value=tot_formula)
        c_sub_val.font = font_grand_total
        c_sub_val.alignment = align_right
        c_sub_val.number_format = '#,##0'

        for c in range(1, 10):
            ws.cell(row=current_row, column=c).border = border_all
        ws.row_dimensions[current_row].height = 24
        before_vat_row = current_row
        current_row += 1

        # Dynamic VAT calculation & display
        effective_vat_rate = (float(vat_percent) / 100.0) if vat_percent is not None else settings.DEFAULT_VAT_RATE
        vat_display_pct = int(round(effective_vat_rate * 100)) if (effective_vat_rate * 100).is_integer() else round(effective_vat_rate * 100, 1)

        # Row: THUẾ GTGT (Dynamic %)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c_vat = ws.cell(row=current_row, column=1, value=f"THUẾ GTGT {vat_display_pct}%")
        c_vat.font = font_grand_total
        c_vat.alignment = align_right

        vat_formula = f"=H{before_vat_row}*{effective_vat_rate}"
        c_vat_val = ws.cell(row=current_row, column=8, value=vat_formula)
        c_vat_val.font = font_grand_total
        c_vat_val.alignment = align_right
        c_vat_val.number_format = '#,##0'

        for c in range(1, 10):
            ws.cell(row=current_row, column=c).border = border_all
        ws.row_dimensions[current_row].height = 24
        vat_row = current_row
        current_row += 1

        # Row: TỔNG GIÁ TRỊ SAU THUẾ (Gold Yellow #FFC000)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=7)
        c_tot = ws.cell(row=current_row, column=1, value="TỔNG GIÁ TRỊ SAU THUẾ")
        c_tot.font = font_grand_total
        c_tot.alignment = align_right

        grand_tot_formula = f"=H{before_vat_row}+H{vat_row}"
        c_tot_val = ws.cell(row=current_row, column=8, value=grand_tot_formula)
        c_tot_val.font = font_grand_total
        c_tot_val.alignment = align_right
        c_tot_val.number_format = '#,##0'

        for c in range(1, 10):
            ws.cell(row=current_row, column=c).border = border_all
            ws.cell(row=current_row, column=c).fill = fill_total
        ws.row_dimensions[current_row].height = 26
        current_row += 1

        # 5. Dedicated Technical Proposals Table (BẢNG ĐỀ XUẤT PHƯƠNG ÁN KỸ THUẬT & THIẾT BỊ TƯƠNG THÍCH DO AI PHÂN TÍCH)
        if proposals is None:
            proposals = []
            for d in devices:
                if isinstance(d, dict) and (d.get("compatible_proposal") or d.get("is_alternative_recommended")):
                    cp = d.get("compatible_proposal") or {}
                    proposals.append({
                        "original_device": cp.get("original_device") or d.get("name") or "",
                        "original_spec": cp.get("original_spec") or d.get("spec") or "",
                        "ai_analysis": cp.get("ai_analysis") or d.get("notes") or "",
                        "proposed_device": cp.get("proposed_device") or d.get("sku") or "",
                        "proposed_spec": cp.get("proposed_spec") or d.get("spec") or "",
                        "suggested_brand": cp.get("suggested_brand") or d.get("origin") or "",
                        "technical_reason": cp.get("technical_reason") or d.get("compatibility_note") or "Bảo toàn 100% sơ đồ nguyên lý"
                    })

        if proposals:
            current_row += 1  # Dòng trống phân cách

            # Title banner
            fill_prop_title = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
            font_prop_title = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
            c_prop_title = ws.cell(row=current_row, column=1, value="BẢNG ĐỀ XUẤT PHƯƠNG ÁN KỸ THUẬT & THIẾT BỊ TƯƠNG THÍCH (DO AI PHÂN TÍCH)")
            c_prop_title.font = font_prop_title
            c_prop_title.fill = fill_prop_title
            c_prop_title.alignment = align_center
            for c in range(1, 10):
                ws.cell(row=current_row, column=c).border = border_all
                ws.cell(row=current_row, column=c).fill = fill_prop_title
            ws.row_dimensions[current_row].height = 28
            current_row += 1

            # Subtitle commitment banner
            fill_prop_sub = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
            font_prop_sub = Font(name="Times New Roman", size=10, italic=True, bold=True, color="1F4E79")
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=9)
            c_sub_note = ws.cell(row=current_row, column=1, value="Cam kết kỹ thuật: Bảo toàn 100% sơ đồ nguyên lý 1 sợi (SLD), số cực, dòng ngắn mạch Icu và tính chọn lọc bảo vệ.")
            c_sub_note.font = font_prop_sub
            c_sub_note.fill = fill_prop_sub
            c_sub_note.alignment = align_center
            for c in range(1, 10):
                ws.cell(row=current_row, column=c).border = border_all
                ws.cell(row=current_row, column=c).fill = fill_prop_sub
            ws.row_dimensions[current_row].height = 22
            current_row += 1

            # Table Header for Proposals
            fill_prop_head = PatternFill(start_color="8EA9DB", end_color="8EA9DB", fill_type="solid")
            font_prop_head = Font(name="Times New Roman", size=10, bold=True, color="000000")

            ws.cell(row=current_row, column=1, value="TT").font = font_prop_head
            ws.cell(row=current_row, column=2, value="THIẾT BỊ GỐC (BẢN VẼ)").font = font_prop_head
            ws.cell(row=current_row, column=3, value="QUY CÁCH GỐC").font = font_prop_head
            ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
            ws.cell(row=current_row, column=4, value="PHÂN TÍCH KỸ THUẬT AI").font = font_prop_head
            ws.merge_cells(start_row=current_row, start_column=6, end_row=current_row, end_column=7)
            ws.cell(row=current_row, column=6, value="THIẾT BỊ ĐỀ XUẤT TƯƠNG THÍCH").font = font_prop_head
            ws.cell(row=current_row, column=8, value="HÃNG SX").font = font_prop_head
            ws.cell(row=current_row, column=9, value="CAM KẾT SLD").font = font_prop_head

            for c in range(1, 10):
                ws.cell(row=current_row, column=c).border = border_all
                ws.cell(row=current_row, column=c).fill = fill_prop_head
                ws.cell(row=current_row, column=c).alignment = align_center
            ws.row_dimensions[current_row].height = 26
            current_row += 1

            # Data rows
            for p_idx, prop in enumerate(proposals, 1):
                ws.cell(row=current_row, column=1, value=str(p_idx)).alignment = align_center
                ws.cell(row=current_row, column=2, value=str(prop.get("original_device") or "")).alignment = align_left
                ws.cell(row=current_row, column=3, value=str(prop.get("original_spec") or "")).alignment = align_center

                ws.merge_cells(start_row=current_row, start_column=4, end_row=current_row, end_column=5)
                ws.cell(row=current_row, column=4, value=str(prop.get("ai_analysis") or "")).alignment = align_left

                ws.merge_cells(start_row=current_row, start_column=6, end_row=current_row, end_column=7)
                prop_dev_text = str(prop.get("proposed_device") or "")
                if prop.get("proposed_spec"):
                    prop_dev_text += f"\n({prop.get('proposed_spec')})"
                ws.cell(row=current_row, column=6, value=prop_dev_text).alignment = align_left

                ws.cell(row=current_row, column=8, value=str(prop.get("suggested_brand") or "")).alignment = align_center
                ws.cell(row=current_row, column=9, value=str(prop.get("technical_reason") or "Bảo toàn 100% SLD")).alignment = align_left

                for c in range(1, 10):
                    ws.cell(row=current_row, column=c).font = font_data
                    ws.cell(row=current_row, column=c).border = border_all
                ws.row_dimensions[current_row].height = 36
                current_row += 1

        # Column widths auto-adjustment
        col_widths = {
            1: 6,   # TT
            2: 45,  # MÔ TẢ CHI TIẾT
            3: 22,  # MÃ HÀNG
            4: 14,  # XUẤT XỨ
            5: 10,  # ĐƠN VỊ
            6: 12,  # SỐ LƯỢNG
            7: 16,  # ĐƠN GIÁ
            8: 18,  # THÀNH TIỀN
            9: 25   # GHI CHÚ
        }
        col_widths.update({10: 18, 11: 20})
        for col_idx, width in col_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

        # 7. Save file to exports directory
        out_dir = Path(settings.EXPORT_DIR)
        out_dir.mkdir(parents=True, exist_ok=True)
        clean_proj = "".join(c for c in project_name if c.isalnum() or c in ("-", "_", " ")).strip().replace(" ", "_")
        clean_suffix = "".join(c for c in (filename_suffix or "") if c.isalnum() or c in ("-", "_", " ")).strip().replace(" ", "_")
        suffix_part = f"_{clean_suffix}" if clean_suffix else ""
        filename = f"BaoGia_{clean_proj}{suffix_part}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:10]}.xlsx"
        file_path = out_dir / filename
        wb.save(file_path)
        return str(file_path)

    @staticmethod
    def match_devices_multi_brand(devices: list, default_brand: str = "Theo thiết kế") -> list:
        """Đối chiếu các thiết bị bóc tách với tất cả các hãng trong Catalog DB."""
        from app.services.device_catalog_engine import DeviceCatalogEngine
        eng = DeviceCatalogEngine()

        brand_mapping = BRAND_CATALOG_MAPPING

        # Gộp các thiết bị trùng lặp trước khi đối chiếu catalog
        devices = DynamicGroupingEngine.merge_duplicates(devices)

        matched_items = []
        for idx, dev in enumerate(devices, 1):
            cat = str(dev.get("category") or "Thiết bị").strip()
            name = str(dev.get("name") or "Thiết bị").strip()
            spec = str(dev.get("spec") or "").strip()
            poles = dev.get("poles")
            in_a = dev.get("in_a")
            qty = int(dev.get("quantity") or 1)

            # Build brand options by querying DeviceCatalogEngine
            brand_options = {}
            for display_name, b_key in brand_mapping.items():
                matches = eng.filter_devices(
                    brand=b_key,
                    device_type=cat,
                    poles=poles,
                    in_current=in_a,
                    limit=1
                )
                if matches:
                    m = matches[0]
                    p_val = int(m.get("g") or m.get("price") or 0)
                    brand_options[display_name] = {
                        "sku": m.get("ma") or m.get("sku") or "",
                        "name": m.get("n") or m.get("name") or name,
                        "price": p_val
                    }

            # Check if device already has a brand from drawing
            dev_b = str(dev.get("brand") or "").strip()
            dev_matched_brand = None
            for b_name in brand_options.keys():
                if dev_b and (b_name.lower() in dev_b.lower() or dev_b.lower() in b_name.lower()):
                    dev_matched_brand = b_name
                    break

            if not brand_options:
                matched_items.append({
                    "id": idx, "name": name, "category": cat, "spec": spec,
                    "poles": poles, "in_a": in_a, "quantity": qty, "unit": "Cái",
                    "selected_brand": dev_b, "sku": "", "description": name,
                    "unit_price": 0, "discount_percent": 0, "line_total": 0,
                    "brand_options": {}, "catalog_matched": False,
                })
                continue

            # If user explicitly chooses a specific brand (e.g. 'Schneider', 'ABB') use that;
            # otherwise prioritize original brand from drawing, or fallback gracefully
            if default_brand and default_brand not in ["Theo thiết kế", "Gốc", "all", ""]:
                active_brand = default_brand if default_brand in brand_options else (dev_matched_brand or list(brand_options.keys())[0])
            else:
                active_brand = dev_matched_brand or list(brand_options.keys())[0]

            selected_opt = brand_options.get(active_brand) or list(brand_options.values())[0]
            final_sku = dev.get("part_number") or selected_opt["sku"]

            matched_items.append({
                "id": idx,
                "name": name,
                "category": cat,
                "spec": spec,
                "poles": poles,
                "in_a": in_a,
                "quantity": qty,
                "unit": "Bộ" if "ACB" in cat or "Tủ" in cat else "Cái",
                "selected_brand": dev_b if (default_brand in ["Theo thiết kế", "Gốc", ""] and dev_b) else active_brand,
                "sku": final_sku,
                "description": selected_opt["name"],
                "unit_price": selected_opt["price"],
                "discount_percent": 0,
                "line_total": qty * selected_opt["price"],
                "brand_options": brand_options
                ,"catalog_matched": True
            })

        return matched_items
