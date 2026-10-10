"""Read an approved source quotation without turning lots or panel totals into devices.

This parser accepts a table only when its labelled columns are unambiguous.
Workbook content supplies facts, never approval to replace the schematic.
"""
from __future__ import annotations
import copy
import math
import re
from pathlib import Path
from .source_reconciliation import normalize_text, quote_spec


def panel_key(value):
    return re.sub(r"[^A-Z0-9]", "", re.sub(r"^TD[- .]*", "", normalize_text(value)))


def category(text):
    n = normalize_text(text)
    direct = quote_spec(text)['category']
    if direct:
        return direct
    if re.search(r'\bTIMER\b', n): return 'TIMER'
    if 'CONG TAC 3 VI TRI' in n: return 'SELECTOR_3_POSITION'
    if 'BO NGUON' in n: return 'POWER_SUPPLY'
    if re.search(r'RO\s*LE|RELAY', n): return 'RELAY'
    return None


def read_configuration(path):
    from openpyxl import load_workbook
    path = Path(path)
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        tables = []
        for sheet in book.worksheets:
            columns = None
            panels, items, other, current, section = [], [], [], None, ''
            for row_no, values in enumerate(sheet.iter_rows(values_only=True), 1):
                labels = {normalize_text(v): i for i, v in enumerate(values) if v is not None}
                header = {}
                for label, idx in labels.items():
                    if label == 'MO TA CHI TIET': header['description'] = idx
                    elif label == 'MA HANG': header['model'] = idx
                    elif 'HANG SX' in label: header['brand'] = idx
                    elif label == 'DON VI': header['unit'] = idx
                    elif label.replace('.', '').replace(' ', '') == 'SLUONG': header['quantity'] = idx
                    elif label.startswith('DON GIA'): header['price'] = idx
                if len(header) >= 5 and all(k in header for k in ('description','model','brand','unit','quantity')):
                    if columns is not None: raise ValueError('Bảng báo giá lặp tiêu đề; cần xác minh không trùng BOM.')
                    columns = header
                    continue
                if columns is None: continue
                def cell(key):
                    i = columns.get(key)
                    return values[i] if i is not None and i < len(values) else None
                text = str(cell('description') or '').strip()
                if not text: continue
                match = re.search(r'\(\s*(T[ĐD][- .]*[A-Z0-9. -]+)\s*\)', text, re.I)
                if match:
                    current = match[1].strip()
                    panels.append({'panel_code': current, 'evidence': f'{path.name} / {sheet.title}!{row_no}',
                                   'package_price_vnd': cell('price'), 'dimension': None})
                    section = ''
                    continue
                if not current: continue
                if normalize_text(text) in ('DAU VAO', 'DAU RA'):
                    section = text
                    continue
                qty = cell('quantity')
                if not isinstance(qty, (int,float)) or isinstance(qty,bool) or not math.isfinite(qty) or qty <= 0:
                    continue
                fact = {**quote_spec(text), 'category': category(text), 'panel_code': current,
                        'quantity': qty, 'model': str(cell('model') or '').strip() or None,
                        'brand': str(cell('brand') or '').strip() or None, 'unit': cell('unit'),
                        'unit_price_vnd': cell('price'), 'section': section,
                        'evidence': f'{path.name} / {sheet.title}!{row_no}', 'row': row_no}
                dim = re.search(r'H\s*(\d+)\s*[x×]\s*W\s*(\d+)\s*[x×]\s*D\s*(\d+)', text, re.I)
                if dim:
                    panels[-1]['dimension'] = f'{dim[2]}x{dim[1]}x{dim[3]}'
                    panels[-1]['enclosure_observation'] = fact
                if fact['category']:
                    if int(qty) != qty: raise ValueError(f"Số lượng thiết bị không nguyên: {fact['evidence']}")
                    fact['quantity'] = int(qty)
                    items.append(fact)
                else:
                    other.append(fact)
            if items: tables.append({'source_filename': path.name, 'panels': panels, 'items': items, 'other_rows': other})
        if len(tables) != 1:
            raise ValueError('Cần đúng một bảng báo giá có cột mô tả, mã hàng, hãng, đơn vị và số lượng để đối chiếu.')
        return tables[0]
    finally:
        book.close()


def schematic_devices(devices):
    source = [copy.deepcopy(d) for d in devices]
    # Re-runs retain original schematic evidence, rather than matching a revised
    # 20A contactor back to the 16A/10A source or adding auxiliaries twice.
    original = []
    for d in source:
        obs = (d.get('source_observations') or {}).get('schematic')
        if d.get('selection_source') == 'approved_quotation':
            if obs: original.append(copy.deepcopy(obs['device']))
        else: original.append(d)
    return original


def apply_configuration(devices, configuration):
    """Bind approved changes to original circuits; never silently drop unmatched circuits."""
    original = schematic_devices(devices)
    panel_codes = {panel_key(d.get('panel_code')): d.get('panel_code') for d in original}
    result, issues = [], []
    remaining = [int(d.get('quantity') or 1) for d in original]
    additions = {'SELECTOR_3_POSITION','POWER_SUPPLY','RELAY'}
    for q in configuration['items']:
        key = panel_key(q['panel_code'])
        if key not in panel_codes: continue  # A per-panel run must not add another panel.
        cat = q['category']
        # The first literal current is the source value before an explicit => revision.
        old_current = re.search(r'(?<![\w.])(\d+(?:[.,]\d+)?)\s*A\b', q['original_text'], re.I)
        current = float(old_current[1].replace(',','.')) if old_current else q.get('current_a')
        matches = [i for i,d in enumerate(original) if remaining[i] and panel_key(d.get('panel_code')) == key
                   and str(d.get('category') or '').upper() == cat
                   and (q.get('poles') is None or (d.get('poles') or quote_spec(d.get('spec') or '')['poles']) == q['poles'])
                   and (current is None or (d.get('in_a') or quote_spec(d.get('spec') or '')['current_a']) == current)]
        needed = q['quantity']
        bound = []
        for i in matches:
            n = min(needed,remaining[i])
            if not n: break
            old = original[i]
            new = copy.deepcopy(old)
            new['quantity'] = n
            source_device = {k:v for k,v in old.items() if k not in (
                'evidence_image','panel_evidence_image','cad','dimensions','catalog_matches','catalog_candidates','catalog_references')}
            source_device['quantity'] = n
            new['source_observations'] = {'schematic': {
                **quote_spec(old.get('spec') or ''), 'device': source_device, 'quantity': n,
                'model': old.get('part_number'), 'brand': old.get('detected_brand'),
                'evidence': old.get('source_filename') or old.get('source_type')}, 'quotation': q}
            bound.append(new)
            remaining[i] -= n
            needed -= n
        if needed and cat in additions:
            bound.append({'category': cat, 'name': q['original_text'], 'quantity': needed,
                          'panel_code': panel_codes[key], 'section': 'Điều khiển',
                          'tag': 'QROW' + str(q['row']), 'source_observations': {'quotation': q},
                          'source_filename': configuration['source_filename'], 'source_type': 'xlsx',
                          'confidence': 1.0})
            needed = 0
        if needed:
            issues.append(f"{q['evidence']}: còn {needed} thiết bị chưa ghép với sơ đồ")
        for new in bound:
            spec = ' '.join(str(v) for v in (
                f"{q['poles']}P" if q.get('poles') else None,
                f"{q['current_a']:g}A" if q.get('current_a') else None,
                f"{q['breaking_capacity_ka']:g}kA" if q.get('breaking_capacity_ka') else None) if v)
            new.update(spec=spec or q['original_text'], in_a=q.get('current_a'), icu_ka=q.get('breaking_capacity_ka'),
                       poles=q.get('poles'), brand=q.get('brand'), detected_brand=q.get('brand'),
                       part_number=q.get('model'), selection_source='approved_quotation', cad=None, dimensions=None,
                       original_spec=(new.get('source_observations') or {}).get('schematic', {}).get('device', {}).get('spec'),
                       accompanying_accessories=None, inferred_components=None)
            if cat == 'SELECTOR_3_POSITION': new.update(mounting='DOOR_MOUNTED', mounting_face='inner_door')
            result.append(new)
    issues += [f"{d.get('panel_code')} / {d.get('tag') or d.get('name')}: {remaining[i]} thiết bị sơ đồ chưa có trong báo giá"
               for i,d in enumerate(original) if remaining[i]]
    if issues: raise ValueError('Đối chiếu báo giá chưa đủ: ' + '; '.join(issues))
    return result
